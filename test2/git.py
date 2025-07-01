import json
import base64
from docx import Document
from docx.shared import Pt
from PIL import Image
import io
import os

def extract_paragraph_style(paragraph):
    alignment = str(paragraph.alignment) if paragraph.alignment else "None"
    runs = []
    for run in paragraph.runs:
        font = run.font
        run_info = {
            "text": run.text,
            "bold": font.bold,
            "italic": font.italic,
            "underline": font.underline,
            "font_name": font.name,
            "font_size": font.size.pt if font.size else None,
        }
        runs.append(run_info)
    return {
        "alignment": alignment,
        "runs": runs
    }

def extract_images(docx_path):
    from zipfile import ZipFile

    image_data = []
    try:
        with ZipFile(docx_path, 'r') as docx_zip:
            for file in docx_zip.namelist():
                if file.startswith('word/media/'):
                    img_bytes = docx_zip.read(file)
                    encoded = base64.b64encode(img_bytes).decode('utf-8')
                    ext = os.path.splitext(file)[1].replace('.', '')
                    image_data.append({
                        "filename": file,
                        "format": ext,
                        "base64": encoded
                    })
    except Exception as e:
        print(f"Error reading images: {e}")
    return image_data

def extract_docx_to_json(file_path):
    try:
        document = Document(file_path)
        doc_json = {
            "paragraphs": [],
            "images": extract_images(file_path),
            "tables": []
        }

        for para in document.paragraphs:
            para_data = extract_paragraph_style(para)
            doc_json["paragraphs"].append(para_data)

        for table in document.tables:
            table_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    cell_data = [extract_paragraph_style(p) for p in cell.paragraphs]
                    row_data.append(cell_data)
                table_data.append(row_data)
            doc_json["tables"].append(table_data)

        return doc_json

    except Exception as e:
        return {"error": str(e)}

if __name__ == "__main__":
    import sys
    if len(sys.argv) != 2:
        print("Usage: python docx_to_json.py <file_path>")
    else:
        file_path = sys.argv[1]
        result = extract_docx_to_json(file_path)
        with open("output.json", "w", encoding="utf-8") as f:
            json.dump(result, f, indent=2, ensure_ascii=False)
        print("✅ JSON exported to output.json")
